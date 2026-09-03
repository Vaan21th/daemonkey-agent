"""原位改 pptx/docx/xlsx 里圈到的字。复制原件再换文本，页数/图片/底不动。"""
from __future__ import annotations

import shutil
from pathlib import Path

from workers.office_revise import fold_key, fold_span

_PATCH = {".pptx", ".docx", ".xlsx"}


def looks_like_page_md(body: str) -> bool:
    t = (body or "").strip()
    if not t:
        return False
    if "<!--" in t or "\n---\n" in t or t.startswith("---"):
        return True
    lines = [ln for ln in t.splitlines() if ln.strip()]
    heads = sum(1 for ln in lines if ln.lstrip().startswith("#"))
    return heads >= 2 or (len(lines) >= 4 and t.lstrip().startswith("#"))


def can_patch(src: Path, excerpt: str, body: str) -> bool:
    if src.suffix.lower() not in _PATCH:
        return False
    if not (excerpt or "").strip():
        return False
    return not looks_like_page_md(body)


def _para_text(para) -> str:
    runs = list(getattr(para, "runs", []) or [])
    if runs:
        return "".join(r.text or "" for r in runs)
    return getattr(para, "text", "") or ""


def _set_para(para, text: str) -> None:
    runs = list(getattr(para, "runs", []) or [])
    if not runs:
        para.text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def _score(full: str, excerpt: str) -> int:
    needle = (excerpt or "").strip()
    if not needle or not (full or "").strip():
        return 0
    fk, fn = fold_key(full), fold_key(needle)
    if fk and fn and fk == fn:
        return 4
    if needle in full and full.count(needle) == 1:
        return 3
    if fold_span(full, needle):
        return 3
    if fk and fn and (fn.startswith(fk) or fk.startswith(fn)):
        return 1
    return 0


def _apply(para, excerpt: str, new: str) -> None:
    full = _para_text(para)
    needle = excerpt.strip()
    if needle in full and full.count(needle) == 1:
        _set_para(para, full.replace(needle, new, 1))
        return
    span = fold_span(full, needle)
    if span:
        _set_para(para, full[: span[0]] + new + full[span[1] :])
        return
    fk, fn = fold_key(full), fold_key(needle)
    if fk and fn and fn.startswith(fk) and fn != fk and "·" in needle:
        _set_para(para, new.split("·")[0].strip() if "·" in new else new)
        return
    _set_para(para, new)


def _pick(cands: list, excerpt: str, page: int | None, label: str):
    if page:
        cands = [c for c in cands if c[0] == page]
        if not cands:
            raise ValueError(f"第{page}{label}找不到「{excerpt.strip()[:40]}」")
    ranked: list[tuple[int, tuple]] = []
    for c in cands:
        sc = _score(_para_text(c[-1]), excerpt)
        if sc:
            ranked.append((sc, c))
    if not ranked:
        where = f"第{page}{label}" if page else "成品"
        raise ValueError(f"{where}里找不到「{excerpt.strip()[:40]}」")
    top = max(r[0] for r in ranked)
    hits = [c for sc, c in ranked if sc == top]
    if len(hits) > 1:
        raise ValueError(
            f"「{excerpt.strip()[:40]}」对上 {len(hits)} 处，加上 page 收窄"
        )
    return hits[0]


def _iter_frames(shape):
    if getattr(shape, "has_text_frame", False):
        yield shape.text_frame
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame
    try:
        for inner in shape.shapes:
            yield from _iter_frames(inner)
    except Exception:
        return


def _count_pics(prs) -> int:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    n = 0
    for s in prs.slides:
        for sh in s.shapes:
            try:
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    n += 1
            except Exception:
                pass
    return n


def _patch_pptx(src: Path, dest: Path, excerpt: str, new: str, page: int | None) -> dict:
    from pptx import Presentation

    shutil.copy2(src, dest)
    prs = Presentation(str(dest))
    before = len(prs.slides)
    pics = _count_pics(prs)
    cands = []
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            for tf in _iter_frames(shape):
                for para in tf.paragraphs:
                    cands.append((i, para))
    _, para = _pick(cands, excerpt, page, "页")
    _apply(para, excerpt, new)
    prs.save(str(dest))
    after = len(Presentation(str(dest)).slides)
    if after != before:
        dest.unlink(missing_ok=True)
        raise RuntimeError("原位改把页数改了，已撤回")
    return {"replaced": 1, "pages": after, "pics": pics}


def _iter_docx_paras(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for sec in doc.sections:
        yield from sec.header.paragraphs
        yield from sec.footer.paragraphs


def _patch_docx(src: Path, dest: Path, excerpt: str, new: str, page: int | None) -> dict:
    from docx import Document

    shutil.copy2(src, dest)
    doc = Document(str(dest))
    cands = [(1, p) for p in _iter_docx_paras(doc)]
    _, para = _pick(cands, excerpt, None, "篇")
    _apply(para, excerpt, new)
    doc.save(str(dest))
    return {"replaced": 1, "pages": 1, "pics": 0}


def _patch_xlsx(src: Path, dest: Path, excerpt: str, new: str, page: int | None) -> dict:
    from openpyxl import load_workbook

    shutil.copy2(src, dest)
    wb = load_workbook(dest)
    sheets = list(wb.worksheets)
    if page:
        if page < 1 or page > len(sheets):
            dest.unlink(missing_ok=True)
            raise ValueError(f"没有第{page}张表")
        sheets = [sheets[page - 1]]
    needle = excerpt.strip()
    hits = []
    for ws in sheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str) and needle and needle in val:
                    hits.append(cell)
    if len(hits) != 1:
        dest.unlink(missing_ok=True)
        raise ValueError(f"表里「{needle[:40]}」对上 {len(hits)} 格，加上 page 或圈更长")
    hits[0].value = hits[0].value.replace(needle, new, 1)
    wb.save(dest)
    return {"replaced": 1, "pages": len(wb.worksheets), "pics": 0}


def patch_office(
    src: Path,
    dest: Path,
    *,
    excerpt: str,
    new: str,
    page: int | None = None,
) -> dict:
    if src.suffix.lower() not in _PATCH:
        raise ValueError(f"原位改不支持 {src.suffix}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.resolve() == src.resolve():
        raise ValueError("原位改不能盖掉原文")
    new = (new or "").strip()
    if not new:
        raise ValueError("body 是空的")
    if not (excerpt or "").strip():
        raise ValueError("excerpt 是空的")
    ext = src.suffix.lower()
    if ext == ".pptx":
        return _patch_pptx(src, dest, excerpt, new, page)
    if ext == ".docx":
        return _patch_docx(src, dest, excerpt, new, page)
    return _patch_xlsx(src, dest, excerpt, new, page)
