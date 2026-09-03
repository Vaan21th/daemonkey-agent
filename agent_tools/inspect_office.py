"""AUTO · 回看已落盘的 Word / PPT / Excel（文本优先，有图再 look_at）。"""
from __future__ import annotations

from pathlib import Path

from . import TIER_AUTO, ToolResult, ToolSpec, register_tool

_ROOT = Path(__file__).resolve().parent.parent
_ALLOW = (
    _ROOT / "data" / "reports",
    _ROOT / "data" / "presentations",
    _ROOT / "data" / "spreadsheets",
)
_OK_EXT = {".docx", ".pptx", ".xlsx"}


def _summarize(args: dict) -> str:
    return f"回看办公文件 · {(args.get('path') or '').strip() or '?'}"


def _resolve(rel: str) -> Path | ToolResult:
    from daemon_runtime import RUNTIME
    from workers.session_docs import accept_upload
    path, err = accept_upload(getattr(RUNTIME, "session_id", "") or "", rel, root=_ROOT)
    if path is None:
        return ToolResult(ok=False, output="", error=err or "只读办公稿")
    if path.suffix.lower() not in _OK_EXT:
        return ToolResult(ok=False, output="", error="只看 docx / pptx / xlsx")
    return path


def _extract_xlsx(path: Path, limit: int = 30) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=False, read_only=True)
    chunks = []
    for name in wb.sheetnames:
        ws = wb[name]
        chunks.append(f"# {name}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= limit:
                chunks.append("…")
                break
            chunks.append(" | ".join("" if c is None else str(c) for c in row))
    wb.close()
    return "\n".join(chunks)


def _extract_docx(path: Path, limit: int = 80) -> str:
    from docx import Document

    lines = []
    for p in Document(str(path)).paragraphs:
        t = (p.text or "").strip()
        if t:
            lines.append(t)
        if len(lines) >= limit:
            lines.append("…")
            break
    return "\n".join(lines) or "(空文档)"


def _extract_pptx(path: Path, limit: int = 20) -> str:
    from pptx import Presentation

    lines = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        if i > limit:
            lines.append("…")
            break
        texts = []
        charts = []
        for sh in slide.shapes:
            if getattr(sh, "has_chart", False):
                try:
                    charts.append(str(sh.chart.chart_type).split(".")[-1])
                except Exception:
                    charts.append("chart")
            if getattr(sh, "has_text_frame", False):
                t = (sh.text_frame.text or "").strip()
                if t:
                    texts.append(t)
        lines.append(f"# 第 {i} 页")
        if charts:
            lines.append("图表: " + ", ".join(charts))
        lines.extend(texts[:12])
    return "\n".join(lines) or "(空演示稿)"


def _fallback_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".xlsx":
        return _extract_xlsx(path)
    if ext == ".docx":
        return _extract_docx(path)
    return _extract_pptx(path)


def _has_eq_formula(text: str) -> bool:
    for line in text.splitlines():
        for part in line.replace("\t", "|").split("|"):
            if part.strip().startswith("="):
                return True
    return False


def _officecli_text(path: Path) -> str:
    try:
        from workers import officecli
        if officecli.available():
            return (officecli.view_text(path) or "").strip()
    except Exception:
        return ""
    return ""


def _read_body(path: Path) -> tuple[str, str]:
    cli = _officecli_text(path)
    if path.suffix.lower() == ".xlsx":
        try:
            formulas = _extract_xlsx(path)
        except Exception:
            formulas = ""
        # view text 只给算完的数；改稿要看见 = 原文
        if formulas.strip() and cli and _has_eq_formula(formulas):
            return f"{formulas}\n\n# 计算后\n{cli}", "officecli+formula"
        if cli:
            return cli, "officecli"
        if formulas.strip():
            return formulas, "fallback"
        raise RuntimeError("empty xlsx")
    if cli:
        return cli, "officecli"
    return _fallback_text(path), "fallback"


def _run(args: dict) -> ToolResult:
    got = _resolve(str(args.get("path") or ""))
    if isinstance(got, ToolResult):
        return got
    path = got
    try:
        text, engine = _read_body(path)
    except Exception as e:
        return ToolResult(ok=False, output="", error=f"读不开: {type(e).__name__}: {e}")
    if not text.strip():
        return ToolResult(ok=False, output="", error="读不开: 空内容")

    rel = path.relative_to(_ROOT).as_posix()
    lines = [
        f"回看 {rel} · 引擎 {engine}",
        "xlsx 先给公式原文；有 OfficeCLI 再附计算后。",
        "",
        text[:6000],
    ]
    if engine.startswith("officecli"):
        lines.append("")
        lines.append("若要看版式，产物库点「成品」；或 look_at 预览 PNG。")
    return ToolResult(ok=True, output="\n".join(lines))


SPEC = ToolSpec(
    name="inspect_office",
    description=(
        "回看已生成的 docx/pptx/xlsx 文本（公式、表头、页大纲）。"
        "出稿后核对用。圈字改句直接 revise_office；加图/换版式前可以先回看这一页。"
        "只读 data/reports|presentations|spreadsheets。"
    ),
    tier=TIER_AUTO,
    input_schema={
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "相对工程根的文件路径，例如 data/spreadsheets/台账.xlsx",
            },
        },
        "required": ["path"],
        "additionalProperties": False,
    },
    run=_run,
    summarize=_summarize,
)
register_tool(SPEC)
