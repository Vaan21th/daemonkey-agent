"""openpyxl 写出可打开的 xlsx · 表头冻结 + 公式原样写入。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="6B46C1")
_ALT_FILL = PatternFill("solid", fgColor="F5F3FF")
_WHITE = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
_BODY = Font(name="微软雅黑", size=11)
_THIN = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)


def _unique_names(sheets: list[dict]) -> list[str]:
    seen: dict[str, int] = {}
    names = []
    for i, sh in enumerate(sheets, 1):
        base = (str(sh.get("name") or f"Sheet{i}"))[:31] or f"Sheet{i}"
        n = seen.get(base, 0)
        seen[base] = n + 1
        names.append(base if n == 0 else f"{base[:28]}_{n+1}")
    return names


def _write_sheet(ws, headers: list, rows: list[list]) -> None:
    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col, title)
        cell.fill = _HEADER_FILL
        cell.font = _WHITE
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN
    for r_i, row in enumerate(rows, 2):
        alt = r_i % 2 == 0
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, val)
            cell.font = _BODY
            cell.border = _THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if alt:
                cell.fill = _ALT_FILL
            if isinstance(val, float):
                cell.number_format = "0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    ws.row_dimensions[1].height = 22
    for col, title in enumerate(headers, 1):
        longest = len(str(title))
        for row in rows:
            if col - 1 < len(row):
                longest = max(longest, min(40, len(str(row[col - 1]))))
        ws.column_dimensions[get_column_letter(col)].width = min(40, max(10, longest + 3))


def render_workbook(sheets: list[dict], output_path: Path) -> Path:
    if not sheets:
        raise ValueError("至少要有一张有表头的表")
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    names = _unique_names(sheets)
    first = True
    for name, spec in zip(names, sheets):
        ws = wb.active if first else wb.create_sheet(title=name)
        if first:
            ws.title = name
            first = False
        headers = spec.get("headers") or []
        rows = spec.get("rows") or []
        if not headers:
            raise ValueError(f"工作表 {name!r} 没有表头")
        _write_sheet(ws, headers, rows)
    try:
        wb.save(str(output_path))
    except PermissionError:
        stem, suf = output_path.stem, output_path.suffix
        for n in range(1, 8):
            alt = output_path.with_name(f"{stem}__{n}{suf}")
            try:
                wb.save(str(alt))
                return alt
            except PermissionError:
                continue
        raise
    return output_path
