"""pptx 页级手术：抄页、换一页、按页插入、相对坐标贴图。"""
from __future__ import annotations

import io
import shutil
from copy import deepcopy
from pathlib import Path

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_A_BLIP = f"{_A}blip"
_SHAPE_TAIL = ("}sp", "}pic", "}grpSp", "}cxnSp", "}graphicFrame")


def _clear_shapes(slide) -> None:
    tree = slide.shapes._spTree
    for el in [c for c in list(tree) if c.tag.endswith(_SHAPE_TAIL)]:
        tree.remove(el)


def _rewire_images(src_slide, dest_slide, el) -> None:
    for blip in el.iter(_A_BLIP):
        rId = blip.get(_R_EMBED)
        if not rId:
            continue
        try:
            part = src_slide.part.related_part(rId)
            blob = part.blob
        except Exception:
            continue
        _part, new_id = dest_slide.part.get_or_add_image_part(io.BytesIO(blob))
        blip.set(_R_EMBED, new_id)


def _copy_slide_bg(src, dest) -> None:
    from pptx.oxml.ns import qn

    src_csld = src._element.find(qn("p:cSld"))
    dest_csld = dest._element.find(qn("p:cSld"))
    if src_csld is None or dest_csld is None:
        return
    src_bg = src_csld.find(qn("p:bg"))
    if src_bg is None or src_bg.find(f".//{_A_BLIP}") is not None:
        return
    dest_bg = dest_csld.find(qn("p:bg"))
    if dest_bg is not None:
        dest_csld.remove(dest_bg)
    dest_csld.insert(0, deepcopy(src_bg))


def paste_slide(src_slide, dest_slide, *, copy_bg: bool = True) -> None:
    _clear_shapes(dest_slide)
    if copy_bg:
        _copy_slide_bg(src_slide, dest_slide)
    tree = dest_slide.shapes._spTree
    for shape in src_slide.shapes:
        el = deepcopy(shape.element)
        _rewire_images(src_slide, dest_slide, el)
        tree.append(el)


def _layout(prs, prefer=None):
    if prefer is not None:
        return prefer.slide_layout
    return prs.slide_layouts[0]


def _move_tail(prs, added: int, after: int) -> None:
    """after：原稿 1-based；0 插到最前。"""
    lst = prs.slides._sldIdLst
    kids = list(lst)
    if added <= 0 or added >= len(kids):
        return
    new_ids = kids[-added:]
    old_n = len(kids) - added
    pos = max(0, min(int(after), old_n))
    if pos >= old_n:
        return
    for el in new_ids:
        lst.remove(el)
    remain = list(lst)
    anchor = remain[pos]
    for el in reversed(new_ids):
        anchor.addprevious(el)


def graft_slides(base: Path, extra: Path, dest: Path, *, after: int | None = None) -> dict:
    """把 extra 各页接到 dest。after=None 接最后；0 最前；N 第 N 页后。"""
    from pptx import Presentation

    shutil.copy2(base, dest)
    dest_prs = Presentation(str(dest))
    extra_prs = Presentation(str(extra))
    base_n = len(dest_prs.slides)
    last = dest_prs.slides[base_n - 1] if base_n else None
    layout = _layout(dest_prs, last)
    added = 0
    for src in extra_prs.slides:
        new = dest_prs.slides.add_slide(layout)
        if last is not None:
            _copy_slide_bg(last, new)
        paste_slide(src, new, copy_bg=False)
        added += 1
    if after is not None and added:
        _move_tail(dest_prs, added, after)
    dest_prs.save(str(dest))
    return {"base_pages": base_n, "added": added, "total": base_n + added, "after": after}


def replace_slide(base: Path, extra: Path, dest: Path, page: int) -> dict:
    """用 extra 第一页覆盖 dest 的第 page 页（1-based），其余页 XML 不动。"""
    from pptx import Presentation

    if page < 1:
        raise ValueError("page 从 1 起")
    shutil.copy2(base, dest)
    dest_prs = Presentation(str(dest))
    extra_prs = Presentation(str(extra))
    n = len(dest_prs.slides)
    if page > n:
        raise ValueError(f"没有第{page}页（共 {n} 页）")
    if not extra_prs.slides:
        raise ValueError("新页是空的")
    paste_slide(extra_prs.slides[0], dest_prs.slides[page - 1], copy_bg=True)
    dest_prs.save(str(dest))
    return {"page": page, "pages": n}


def _clamp01(v, default: float) -> float:
    try:
        x = float(v)
    except (TypeError, ValueError):
        x = default
    return max(0.0, min(1.0, x))


def place_picture(
    src: Path,
    dest: Path,
    image: Path,
    *,
    page: int,
    x: float = 0.5,
    y: float = 0.5,
    w: float = 0.36,
) -> dict:
    """x/y 是页内 0-1，图钉当图心。w 是相对页宽。"""
    from pptx import Presentation
    from pptx.util import Emu

    if not image.is_file() or image.stat().st_size < 32:
        raise ValueError("配图文件无效")
    if page < 1:
        raise ValueError("page 从 1 起")
    shutil.copy2(src, dest)
    prs = Presentation(str(dest))
    n = len(prs.slides)
    if page > n:
        raise ValueError(f"没有第{page}页（共 {n} 页）")
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    rw = max(0.12, min(0.9, _clamp01(w, 0.36)))
    box_w = int(sw * rw)
    rx, ry = _clamp01(x, 0.5), _clamp01(y, 0.5)
    left = int(sw * rx - box_w / 2)
    top = int(sh * ry - box_w / 2)
    left = max(0, min(left, sw - box_w))
    top = max(0, min(top, sh - box_w))
    prs.slides[page - 1].shapes.add_picture(str(image), Emu(left), Emu(top), Emu(box_w))
    prs.save(str(dest))
    return {"page": page, "pages": n, "x": rx, "y": ry, "w": rw}
