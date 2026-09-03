"""行内 markdown · 加能点的链接。"""
from __future__ import annotations

import re
from typing import Optional

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from .themes import Theme

INLINE_RE = re.compile(
    r"(\[([^\]]+)\]\((https?://[^)\s]+|/[^)\s]+)\))"
    r"|(\*\*[^*]+?\*\*)"
    r"|(`[^`]+?`)"
    r"|(\*[^*]+?\*)"
)


def _rgb(t: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*t)


def add_hyperlink(paragraph, text: str, url: str, theme: Theme, *, size: float) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2B6CB0")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), theme.font_cjk)
    rFonts.set(qn("w:eastAsia"), theme.font_cjk)
    rFonts.set(qn("w:hAnsi"), theme.font_cjk)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(rFonts)
    rPr.append(sz)
    rPr.append(color)
    rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline_runs(
    paragraph,
    text: str,
    theme: Theme,
    set_font,
    *,
    base_size: float = 10.5,
    base_color: Optional[RGBColor] = None,
    base_bold: bool = False,
) -> None:
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos: m.start()])
            set_font(r, theme, size=base_size, color=base_color, bold=base_bold)
        if m.group(1):
            url = m.group(3)
            if url.startswith("//") or not (
                url.startswith("http://") or url.startswith("https://")
                or (url.startswith("/") and not url.startswith("//"))
            ):
                r = paragraph.add_run(m.group(0))
                set_font(r, theme, size=base_size, color=base_color, bold=base_bold)
            else:
                add_hyperlink(paragraph, m.group(2), url, theme, size=base_size)
        elif m.group(4):
            r = paragraph.add_run(m.group(4)[2:-2])
            set_font(r, theme, size=base_size, color=base_color, bold=True)
        elif m.group(5):
            r = paragraph.add_run(m.group(5)[1:-1])
            set_font(r, theme, name=theme.font_en, size=base_size - 0.5,
                     color=_rgb(theme.color_code_inline), bold=base_bold)
        elif m.group(6):
            r = paragraph.add_run(m.group(6)[1:-1])
            set_font(r, theme, size=base_size, color=base_color,
                     bold=base_bold, italic=True)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        set_font(r, theme, size=base_size, color=base_color, bold=base_bold)
