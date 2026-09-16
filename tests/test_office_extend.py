"""已有 pptx 加页：原页文字留下，新页接在后面。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from workers.office_extend import append_slides, render_new_pages


def _plain(path: Path, lines: list[str]) -> None:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for text in lines:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(8), Inches(1.2))
        box.text_frame.text = text
    prs.save(str(path))


def _texts(path: Path) -> list[str]:
    prs = Presentation(str(path))
    out = []
    for sl in prs.slides:
        bits = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
        out.append(" ".join(bits))
    return out


def test_append_keeps_original_pages(tmp_path):
    base = tmp_path / "base.pptx"
    extra = tmp_path / "extra.pptx"
    dest = tmp_path / "out.pptx"
    _plain(base, ["原稿第一页专属字", "原稿第二页不动"])
    n = render_new_pages("# 新加的一页\n- 补充要点", extra, style="light_studio")
    assert n >= 1
    info = append_slides(base, extra, dest)
    assert info["base_pages"] == 2
    assert info["added"] == n
    assert info["total"] == 2 + n
    texts = _texts(dest)
    assert any("原稿第一页专属字" in t for t in texts)
    assert any("原稿第二页不动" in t for t in texts)
    blob = "\n".join(texts)
    assert "新加的一页" in blob or "补充要点" in blob


def test_style_from_pptx_uses_slide_color(tmp_path):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from workers.office_extend import style_from_pptx

    path = tmp_path / "brand.pptx"
    prs = Presentation()
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    box = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x4B, 0x78, 0x2B)
    box.line.fill.background()
    prs.save(str(path))
    st = style_from_pptx(path)
    assert st is not None
    assert st.accent == "4B782B"
    assert st.accent != "7C3AED"


def test_style_from_user_original():
    from workers.office_extend import style_from_pptx

    src = Path(__file__).resolve().parent.parent / "data" / "presentations" / "AI_AGENT_新增页.pptx"
    if not src.is_file():
        return
    st = style_from_pptx(src)
    assert st is not None
    assert st.accent != "7C3AED"
    assert st.font_en


def test_extend_gate_locks_generate():
    from workers.office_extend import apply_extend_gate, extend_lock

    msg = "办公稿已挂进本话题 · 路径: data/presentations/a.pptx\n把这个拓展到10P"
    assert extend_lock(msg)
    tools, think, hint = apply_extend_gate(msg, "auto")
    assert tools == {"extend_office", "inspect_office"}
    assert think == "off"
    assert "generate_presentation" in hint
    assert not extend_lock("做一份10页PPT")
    assert not extend_lock("【中栏批注】把这个拓展到10P path=data/presentations/a.pptx")


def test_extend_office_tool(tmp_path, monkeypatch):
    from agent_tools import extend_office as eo

    root = tmp_path
    folder = root / "data" / "presentations"
    folder.mkdir(parents=True)
    src = folder / "客户链路.pptx"
    _plain(src, ["原封面还在"])
    monkeypatch.setattr(eo, "_ROOT", root)
    monkeypatch.setattr(eo, "_ALLOW", folder)
    monkeypatch.setattr("workers.output_versions.Path", Path)
    got = eo._run({
        "path": "data/presentations/客户链路.pptx",
        "body": "# 补一页\n- 只加这页",
    })
    assert got.ok, got.error
    assert "原 1 页留下" in got.output
    texts = _texts(folder / "客户链路.pptx")
    assert any("原封面还在" in t for t in texts)
