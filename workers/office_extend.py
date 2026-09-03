"""已有 pptx 只加新页 · 原页 XML 留下，新页抄原稿配色。"""
from __future__ import annotations

import re
from pathlib import Path

_EXTEND_RE = re.compile(r"拓展到|扩展到|加到\s*\d+\s*页|加几页|扩到\s*\d+")
_HAS_DOC = re.compile(r"办公稿已挂进本话题|data/presentations/|\.pptx", re.I)
_HINT = (
    "\n已有稿加页：只准 extend_office + inspect_office。"
    "body 只写新增页。禁止 generate_presentation / read_scenario。"
    "原页板式和素材必须留下。新页抄原稿配色，不要另起一套风格。"
    "after=页码则插在该页后面。\n"
)
_THEME_KEYS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2")


def extend_lock(message: str) -> bool:
    text = message or ""
    if "【中栏批注】" in text or "中栏画布上批了" in text:
        return False
    return bool(_EXTEND_RE.search(text) and _HAS_DOC.search(text))


def apply_extend_gate(message: str, thinking: str | None) -> tuple[set[str] | None, str | None, str]:
    if not extend_lock(message):
        return None, thinking, ""
    tv = (thinking or "auto").strip().lower() or "auto"
    if tv == "auto":
        thinking = "off"
    return {"extend_office", "inspect_office"}, thinking, _HINT


def _hex6(s: str | None) -> str | None:
    s = (s or "").strip().lstrip("#").upper()
    return s if len(s) == 6 and all(c in "0123456789ABCDEF" for c in s) else None


def _clr_hex(el) -> str | None:
    if el is None:
        return None
    for child in el.iter():
        tag = child.tag
        if tag.endswith("}srgbClr"):
            return _hex6(child.get("val"))
        if tag.endswith("}sysClr"):
            return _hex6(child.get("lastClr"))
    return None


def _theme_part(prs):
    for master in prs.slide_masters:
        for rel in master.part.rels.values():
            if "theme" in (rel.reltype or ""):
                return rel.target_part
    return None


def _freq(xs: list[str]) -> str | None:
    if not xs:
        return None
    counts: dict[str, int] = {}
    for x in xs:
        counts[x] = counts.get(x, 0) + 1
    return max(counts, key=counts.get)


def harvest_theme(path: Path) -> dict:
    """从原稿 theme + 前几页取样配色/字体。"""
    from lxml import etree
    from pptx import Presentation

    prs = Presentation(str(path))
    colors, fonts, fills, inks = {}, {}, [], []
    part = _theme_part(prs)
    if part:
        root = etree.fromstring(part.blob)
        for el in root.iter():
            name = el.tag.split("}")[-1]
            if name in _THEME_KEYS:
                h = _clr_hex(el)
                if h:
                    colors[name] = h
            if name != "majorFont":
                continue
            for kid in el:
                tf = (kid.get("typeface") or "").strip()
                kn = kid.tag.split("}")[-1]
                if kn == "latin" and tf:
                    fonts["en"] = tf
                if kn == "ea" and tf:
                    fonts["cjk"] = tf
    for sl in list(prs.slides)[:5]:
        for sh in sl.shapes:
            try:
                if sh.fill.type is not None:
                    h = _hex6(str(sh.fill.fore_color.rgb))
                    if h and h not in {"FFFFFF", "000000"}:
                        fills.append(h)
            except Exception:
                pass
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        h = _hex6(str(run.font.color.rgb))
                        if h:
                            inks.append(h)
                    except Exception:
                        pass
    return {"colors": colors, "fonts": fonts, "fills": fills, "inks": inks}


def style_from_pptx(path: Path):
    """原稿视觉 → DeckStyle。抄不到返回 None。"""
    from slides_engine.styles import _lum, resolve_style

    try:
        bag = harvest_theme(path)
    except Exception:
        return None
    c = bag["colors"]
    bg = c.get("lt1") or "FFFFFF"
    dk = c.get("dk2") or c.get("dk1") or "141414"
    dark = _lum(bg) < 0.45
    accent = _freq(bag["fills"]) or c.get("accent1")
    if not accent:
        return None
    spec = {
        "bg": bg,
        "bg_alt": c.get("lt2") or bg,
        "ink_title": (c.get("lt1") or "F5F3FF") if dark else dk,
        "ink_body": (c.get("lt2") or "D1D5DB") if dark else (c.get("dk1") or "1F2937"),
        "accent": accent,
        "accent2": c.get("accent2") or accent,
        "is_dark": dark,
    }
    fonts = bag["fonts"]
    if fonts.get("cjk"):
        spec["font_cjk"] = fonts["cjk"]
    if fonts.get("en"):
        spec["font_en"] = fonts["en"]
    return resolve_style("dark_keynote" if dark else "light_studio", spec=spec)


def append_slides(base: Path, extra: Path, dest: Path, *, after: int | None = None) -> dict:
    """把 extra 接到 dest。after=None 最后；0 最前；N 第 N 页后。"""
    from workers.office_slides import graft_slides

    return graft_slides(base, extra, dest, after=after)


def render_new_pages(body: str, extra: Path, *, style: str = "light_studio", here=None, inherit=None) -> int:
    """只渲新增页。有 inherit 就抄那份 pptx 的配色，不加自动封面。"""
    from slides_engine import parse_deck, render_deck, resolve_style

    slides = parse_deck(body or "")
    if not slides:
        raise ValueError("新增页 markdown 解析不到任何页（用 --- 分页）")
    extra.parent.mkdir(parents=True, exist_ok=True)
    st = style_from_pptx(inherit) if inherit else None
    if st is None:
        st = resolve_style(style)
    render_deck(slides, extra, cover=None, style=st, here_dir=here)
    return len(slides)
