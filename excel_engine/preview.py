"""xlsx → 单文件 HTML，浏览器就能看，不绑 Excel。"""
from __future__ import annotations

import html
from pathlib import Path

from openpyxl import load_workbook

_CSS = """
body{margin:0;font:14px/1.5 "Microsoft YaHei","Segoe UI",sans-serif;color:#1a1530;background:#f7f5ff}
nav{display:flex;gap:6px;padding:10px 14px;background:#fff;border-bottom:1px solid #e8e4f5;flex-wrap:wrap}
nav button{border:1px solid #d4c8f0;background:#fff;color:#4c1d95;border-radius:8px;padding:4px 10px;cursor:pointer}
nav button.on{background:#6b46c1;color:#fff;border-color:#6b46c1}
main{padding:16px}
table{border-collapse:collapse;min-width:40%;background:#fff;box-shadow:0 1px 4px #6b46c114}
th,td{border:1px solid #e2e8f0;padding:6px 10px;text-align:left}
th{background:#6b46c1;color:#fff;font-weight:600}
tr:nth-child(even) td{background:#f5f3ff}
.formula{font-family:Consolas,monospace;color:#6b46c1}
"""


def _cell_html(val) -> str:
    if val is None:
        return ""
    text = str(val)
    cls = ' class="formula"' if text.startswith("=") else ""
    return f"<td{cls}>{html.escape(text)}</td>"


def workbook_to_html(path: Path, *, max_rows: int = 200) -> str:
    wb = load_workbook(str(path), data_only=False, read_only=True)
    tabs = []
    bodies = []
    for i, name in enumerate(wb.sheetnames):
        ws = wb[name]
        rows = []
        for r_i, row in enumerate(ws.iter_rows(values_only=True)):
            if r_i >= max_rows + 1:
                break
            rows.append(row)
        if not rows:
            continue
        header = "".join(f"<th>{html.escape(str(c or ''))}</th>" for c in rows[0])
        trs = [f"<tr>{''.join(_cell_html(c) for c in row)}</tr>" for row in rows[1:]]
        hid = "" if i == 0 else " hidden"
        on = " on" if i == 0 else ""
        safe = html.escape(name)
        tabs.append(f'<button type="button" class="tab{on}" data-i="{i}">{safe}</button>')
        bodies.append(
            f'<section data-i="{i}"{hid}><h2>{safe}</h2>'
            f"<table><thead><tr>{header}</tr></thead><tbody>{''.join(trs)}</tbody></table></section>"
        )
    wb.close()
    script = (
        "<script>document.querySelectorAll('nav button').forEach(b=>b.onclick=()=>{"
        "const i=b.getAttribute('data-i');"
        "document.querySelectorAll('nav button').forEach(x=>x.classList.toggle('on',x===b));"
        "document.querySelectorAll('section').forEach(s=>s.hidden=s.getAttribute('data-i')!==i);"
        "});</script>"
    )
    return (
        "<!doctype html><meta charset='utf-8'>"
        f"<title>{html.escape(path.name)}</title><style>{_CSS}</style>"
        f"<nav>{''.join(tabs)}</nav><main>{''.join(bodies)}</main>{script}"
    )


def write_preview(src: Path, dest: Path, *, max_rows: int = 200) -> Path:
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(workbook_to_html(src, max_rows=max_rows), encoding="utf-8")
    return dest
